"""Unit tests for the RAG R1 text pipeline: extraction, chunking and personal-data scan."""

from io import BytesIO

from app.modules.knowledge.chunking import CHUNKING_VERSION, TARGET_CHARS, chunk_sections
from app.modules.knowledge.extraction import Section, extract_sections, page_count
from app.modules.knowledge.pii import scan_sections


def section(*paragraphs: str, heading: str = "Returns policy > Damaged items") -> Section:
    return Section(heading, list(paragraphs), 1, 1)


def test_chunking_keeps_chunks_near_the_target_size() -> None:
    paragraphs = [("word " * 220).strip() for _ in range(6)]
    chunks = chunk_sections([section(*paragraphs)])
    assert len(chunks) >= 2
    assert all(chunk.char_count <= TARGET_CHARS + 50 for chunk in chunks)
    assert all(chunk.blocks for chunk in chunks)


def test_chunking_overlaps_one_paragraph_and_keeps_the_heading() -> None:
    paragraphs = [f"Paragraph {index} " + ("text " * 120) for index in range(4)]
    chunks = chunk_sections([section(*paragraphs)])
    assert len(chunks) >= 2
    # The overlap carries the last paragraph of a chunk into the next one.
    assert chunks[1].blocks[0] == chunks[0].blocks[-1]
    assert all(chunk.heading_path == "Returns policy > Damaged items" for chunk in chunks)
    assert chunks[0].embedding_text.startswith("Returns policy > Damaged items")


def test_tiny_tail_is_merged_into_the_previous_chunk() -> None:
    long_paragraph = "a" * (TARGET_CHARS - 50)
    chunks = chunk_sections([section(long_paragraph, "short tail")])
    assert len(chunks) == 1
    assert "short tail" in chunks[0].text


def test_chunking_is_versioned() -> None:
    assert CHUNKING_VERSION == "1"


def test_markdown_heading_paths_and_blocks() -> None:
    data = (
        b"# Returns policy\n\n"
        b"We accept returns within 14 days.\n\n"
        b"## Damaged items\n\n"
        b"A damaged item can be returned for free.\n"
    )
    sections = extract_sections("policy.md", data)
    assert [section.heading_path for section in sections] == [
        "Returns policy",
        "Returns policy > Damaged items",
    ]
    assert sections[1].blocks == ["A damaged item can be returned for free."]


def test_plain_text_sections() -> None:
    sections = extract_sections("notes.txt", b"First paragraph.\n\nSecond paragraph.\n")
    assert len(sections) == 1
    assert sections[0].blocks == ["First paragraph.", "Second paragraph."]


def test_docx_headings_become_the_heading_path() -> None:
    from docx import Document

    document = Document()
    document.add_heading("Delivery policy", level=1)
    document.add_paragraph("Orders arrive in 3 to 5 days.")
    buffer = BytesIO()
    document.save(buffer)
    sections = extract_sections("delivery.docx", buffer.getvalue())
    assert sections[0].heading_path == "Delivery policy"
    assert sections[0].blocks == ["Orders arrive in 3 to 5 days."]


def test_pdf_pages_become_sections_with_page_numbers() -> None:
    from reportlab.pdfgen import canvas

    buffer = BytesIO()
    pdf = canvas.Canvas(buffer)
    pdf.drawString(72, 720, "Returns are accepted within 14 days.")
    pdf.showPage()
    pdf.save()
    sections = extract_sections("returns.pdf", buffer.getvalue())
    assert sections
    assert sections[0].page_from == 1
    assert page_count(sections) == 1


def test_personal_data_scan_reports_only_category_codes() -> None:
    card = scan_sections([section("My card is 4111 1111 1111 1111.")])
    assert card == ["card_number"]
    assert scan_sections([section("PAN ABCDE1234F here.")]) == ["pan"]
    assert scan_sections([section("Aadhaar 1234 5678 9012 here.")]) == ["aadhaar"]
    assert scan_sections([section("Nothing personal, just policy text.")]) == []


def test_no_text_is_an_empty_section_list() -> None:
    assert extract_sections("blank.txt", b"   \n\n   \n") == []
